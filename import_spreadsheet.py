# -*- coding: utf-8 -*-
"""
스프레드시트(Excel / CSV)에서 수정한 번역을 translations.json 및 게임 패치에 반영하는 도구
사용법:
  python import_spreadsheet.py                             # 기본 OverField_Translation_Tables.xlsx 가져오기
  python import_spreadsheet.py my_edit.xlsx                 # 특정 엑셀 파일 가져오기
  python import_spreadsheet.py csv_spreadsheets/Table_46.csv # 특정 CSV 파일 가져오기
"""

import sys
import os
import json
import re

sys.stdout.reconfigure(encoding='utf-8')

def import_from_file(filepath):
    if not os.path.exists(filepath):
        print(f"❌ 파일을 찾을 수 없습니다: {filepath}")
        return

    print(f"[*] 1. {filepath} 로드 중...")
    updates = {} # uid -> new_ko

    if filepath.endswith('.xlsx'):
        import openpyxl
        wb = openpyxl.load_workbook(filepath, data_only=True)
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            # 헤더 확인
            header = [cell.value for cell in ws[1]]
            try:
                id_col = header.index("식별자") + 1
                ko_col = header.index("현재 번역") + 1
            except ValueError:
                continue

            for row in range(2, ws.max_row + 1):
                raw_id = ws.cell(row=row, column=id_col).value
                raw_ko = ws.cell(row=row, column=ko_col).value
                if raw_id:
                    uid = str(raw_id).split()[0].strip()
                    ko_str = str(raw_ko) if raw_ko is not None else ""
                    updates[uid] = ko_str

    elif filepath.endswith('.csv'):
        import csv
        with open(filepath, 'r', encoding='utf-8-sig') as f:
            reader = csv.reader(f)
            header = next(reader, None)
            if not header:
                return
            id_col = header.index("식별자")
            ko_col = header.index("현재 번역")
            for row in reader:
                if len(row) > max(id_col, ko_col):
                    raw_id = row[id_col]
                    raw_ko = row[ko_col]
                    uid = str(raw_id).split()[0].strip()
                    updates[uid] = str(raw_ko)

    print(f"[+] 스프레드시트에서 {len(updates):,}개 항목 읽기 완료")

    print("[*] 2. translations.json 업데이트 중...")
    with open('translations.json', 'r', encoding='utf-8') as f:
        translations = json.load(f)

    changed_count = 0
    for it in translations:
        uid = it.get('uid')
        if uid in updates:
            new_ko = updates[uid]
            if it.get('ko') != new_ko:
                it['ko'] = new_ko
                it['modified'] = True
                changed_count += 1

    with open('translations.json', 'w', encoding='utf-8') as f:
        json.dump(translations, f, indent=2, ensure_ascii=False)

    print(f"[+] translations.json 갱신 완료: 총 {changed_count:,}개 문장 수정 반영됨")
    
    if changed_count > 0:
        print("[*] 3. 게임 패치 바이너리 자동 빌드 실행...")
        import apply_t7_t17_retranslation
        print("[+] 게임 패치 적용 및 동기화 완료!")
    else:
        print("ℹ️ 변경된 문장이 없습니다.")

if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else "OverField_Translation_Tables.xlsx"
    import_from_file(target)
