# -*- coding: utf-8 -*-
"""
OverField Table 1, 7, 17, 20, 46, 63 스프레드시트(.xlsx) 생성 스크립트
열 구성: 식별자 / 화자 / 원문 / 현재 번역
"""

import json
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

print("[*] 1. translations.json 로드 중...")
with open('translations.json', 'r', encoding='utf-8') as f:
    translations = json.load(f)

target_tables = [
    (46, "Table 46 (스토리·대화)"),
    (1,  "Table 1 (아이템·도감)"),
    (7,  "Table 7 (스킬·버프)"),
    (17, "Table 17 (장비·특성)"),
    (20, "Table 20 (편지·공지)"),
    (63, "Table 63 (미니게임)")
]

wb = openpyxl.Workbook()
# 기본 시트 제거
wb.remove(wb.active)

# 스타일 정의
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="Malgun Gothic", size=11, bold=True, color="FFFFFF")
header_align = Alignment(horizontal="center", vertical="center")

data_font = Font(name="Malgun Gothic", size=10)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_center = Alignment(horizontal="center", vertical="center")

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

all_rows_for_combined = []

for t_id, sheet_title in target_tables:
    print(f"[*] {sheet_title} 시트 생성 중...")
    ws = wb.create_sheet(title=sheet_title)
    
    # 헤더 작성
    headers = ["식별자", "화자", "원문", "현재 번역"]
    ws.append(headers)
    ws.row_dimensions[1].height = 26
    
    for col_num in range(1, 5):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    # 데이터 필터링 (원문이나 번역 중 하나라도 내용이 있는 항목)
    t_items = [it for it in translations if it.get('table_id') == t_id and (it.get('ja', '').strip() or it.get('ko', '').strip())]

    for it in t_items:
        uid = it.get('uid', '')
        key = it.get('key', '')
        # 식별자: uid와 key 결합 (key가 있으면 uid (key) 형태)
        identifier = f"{uid} ({key})" if key else uid
        
        # 화자: 한국어 화자 우선, 일본어 화자 병기
        spk_ko = it.get('speaker_ko', '').strip()
        spk_ja = it.get('speaker_ja', '').strip()
        if spk_ko and spk_ja and spk_ko != spk_ja:
            speaker = f"{spk_ko} ({spk_ja})"
        elif spk_ko:
            speaker = spk_ko
        elif spk_ja:
            speaker = spk_ja
        else:
            speaker = ""

        ja_text = it.get('ja', '')
        ko_text = it.get('ko', '')

        row_data = [identifier, speaker, ja_text, ko_text]
        ws.append(row_data)
        all_rows_for_combined.append((f"Table {t_id}", row_data))

    # 열 너비 설정
    ws.column_dimensions['A'].width = 24  # 식별자
    ws.column_dimensions['B'].width = 18  # 화자
    ws.column_dimensions['C'].width = 50  # 원문
    ws.column_dimensions['D'].width = 50  # 현재 번역

    # 첫 행 고정 및 자동 필터
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:D{len(t_items) + 1}"
    print(f"    - {len(t_items):,}개 행 작성 완료")

# 전체 통합 시트 생성 (맨 앞 배치)
print("[*] '전체 통합 (All)' 시트 생성 중...")
ws_all = wb.create_sheet(title="전체 통합 (All)", index=0)
headers_all = ["테이블", "식별자", "화자", "원문", "현재 번역"]
ws_all.append(headers_all)
ws_all.row_dimensions[1].height = 26

for col_num in range(1, 6):
    cell = ws_all.cell(row=1, column=col_num)
    cell.fill = PatternFill(start_color="112F4E", end_color="112F4E", fill_type="solid")
    cell.font = header_font
    cell.alignment = header_align

for tbl_tag, r_data in all_rows_for_combined:
    ws_all.append([tbl_tag] + r_data)

ws_all.column_dimensions['A'].width = 14  # 테이블
ws_all.column_dimensions['B'].width = 24  # 식별자
ws_all.column_dimensions['C'].width = 18  # 화자
ws_all.column_dimensions['D'].width = 48  # 원문
ws_all.column_dimensions['E'].width = 48  # 현재 번역

ws_all.freeze_panes = "A2"
ws_all.auto_filter.ref = f"A1:E{len(all_rows_for_combined) + 1}"
print(f"    - 전체 통합 {len(all_rows_for_combined):,}개 행 작성 완료")

output_file = "OverField_Translation_Tables.xlsx"
print(f"[*] 엑셀 파일 저장 중: {output_file}...")
wb.save(output_file)
print(f"[+] 성공적으로 생성되었습니다! 크기: {open(output_file, 'rb').seek(0, 2):,} bytes")
